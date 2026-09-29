"""
Reporting tests.

A report is only useful if it is accurate and if all three formats describe the
same thing. Two things are pinned here:

1. Correctness -- the numbers in a report come from the same queries the UI
   reads, so a report cannot quietly disagree with the dashboard.
2. Format agreement -- JSON, XLSX, and TXT are three renderings of one
   document. A format that drops a section is worse than no format, because it
   looks complete.

The XLSX is checked against a real spreadsheet reader's expectations without
adding a dependency: the zip must be intact, every part must parse as XML, and
the relationships must resolve. openpyxl is used when it happens to be
installed, and skipped when it is not -- the project has no dependencies by
design and a test must not change that.
"""

from __future__ import annotations

import io
import json
import unittest
import zipfile
import xml.dom.minidom as minidom
from datetime import datetime, timedelta, timezone

from tiox import reporting
from tiox.store.control import ControlPlane

BASE = datetime(2026, 9, 29, 12, 0, tzinfo=timezone.utc)


def make_store() -> ControlPlane:
    store = ControlPlane()
    store.init_schema()

    ep = store.upsert_endpoint({
        "hostname": "wks-01", "ip": "10.0.0.5", "os": "Linux", "agent_key": "k1",
    })
    ep2 = store.upsert_endpoint({
        "hostname": "srv-02", "ip": "10.0.0.6", "os": "Linux", "agent_key": "k2",
    })

    # A completed scan with one finding, so the scan report has content.
    scan = store.start_scan("/srv", "local", "srv-02")
    store.record_finding(scan["scan_id"], {
        "file_path": "/srv/evil.exe",
        "file_hash": "aa" * 32,
        "family": "TestFamily",
        "rule_id": "builtin.hash_exact",
        "severity": "critical",
        "techniques": ["T1486"],
        "event_id": "ev-1",
    })
    store.finish_scan(scan["scan_id"], "completed", files_scanned=42,
                      dirs_scanned=3, threats_found=1, errors=0)
    store.scan_id_for_test = scan["scan_id"]  # type: ignore[attr-defined]
    store.ep1 = ep  # type: ignore[attr-defined]
    return store


class TestReportKinds(unittest.TestCase):
    def setUp(self):
        self.store = make_store()
        self.sid = self.store.scan_id_for_test  # type: ignore[attr-defined]

    def tearDown(self):
        self.store.close()

    def test_every_declared_kind_builds(self):
        """A kind advertised to the UI must actually build.

        The UI offers whatever REPORT_KINDS lists, so a kind that raises here
        would be a menu item that always fails.
        """
        for kind in reporting.REPORT_KINDS:
            with self.subTest(kind=kind):
                doc = reporting.build_document(
                    kind, self.store, window="all", scan_id=self.sid)
                self.assertTrue(doc["sections"], f"{kind} produced no sections")
                self.assertEqual(doc["report"]["kind"], kind)

    def test_scan_kind_requires_a_scan_id(self):
        with self.assertRaises(ValueError):
            reporting.build_document("scan", self.store, window="all")

    def test_scan_kind_rejects_an_unknown_scan(self):
        with self.assertRaises(ValueError):
            reporting.build_document("scan", self.store, window="all",
                                     scan_id="does-not-exist")

    def test_unknown_kind_is_rejected_with_the_valid_list(self):
        with self.assertRaises(ValueError) as ctx:
            reporting.build_document("nonsense", self.store)
        self.assertIn("summary", str(ctx.exception))

    def test_every_report_carries_the_summary_block(self):
        """Context is what makes the detail rows interpretable."""
        for kind in reporting.REPORT_KINDS:
            with self.subTest(kind=kind):
                doc = reporting.build_document(
                    kind, self.store, window="all", scan_id=self.sid)
                self.assertEqual(doc["section_names"][0], "summary")

    def test_scan_report_includes_its_findings(self):
        doc = reporting.build_document("scan", self.store, window="all",
                                       scan_id=self.sid)
        names = doc["section_names"]
        self.assertIn("findings", names)
        findings = next(s for s in doc["sections"] if s["name"] == "findings")
        self.assertEqual(len(findings["rows"]), 1)
        self.assertEqual(findings["rows"][0]["family"], "TestFamily")

    def test_summary_agrees_with_the_store(self):
        """The report must use the same numbers the dashboard shows."""
        doc = reporting.build_document("summary", self.store, window="all")
        summary = next(s for s in doc["sections"] if s["name"] == "summary")
        as_map = {r["metric"]: r["value"] for r in summary["rows"]}
        self.assertEqual(as_map["endpoints_registered"],
                         len(self.store.list_endpoints()))
        self.assertEqual(as_map["scans_run"], self.store.scan_stats()["total"])
        self.assertEqual(as_map["threats_found"], self.store.scan_stats()["threats"])


class TestWindowPhrasing(unittest.TestCase):
    def test_all_time_does_not_read_as_last_all(self):
        """The header is trusted text. "last all" makes the report look broken."""
        self.assertEqual(reporting.window_phrase("all"), "all time")
        self.assertEqual(reporting.window_phrase(""), "all time")
        self.assertNotIn("last all", reporting.to_txt(
            reporting.build_document("summary", make_store(), window="all")
        )[0].decode())

    def test_shorthand_and_iso(self):
        self.assertEqual(reporting.window_phrase("24h"), "the last 24h")
        self.assertEqual(reporting.window_phrase("7d"), "the last 7d")
        self.assertEqual(reporting.window_phrase("2026-01-01"),
                         "since 2026-01-01")


class TestFormats(unittest.TestCase):
    def setUp(self):
        self.store = make_store()
        self.sid = self.store.scan_id_for_test  # type: ignore[attr-defined]
        self.doc = reporting.build_document("summary", self.store, window="all")

    def tearDown(self):
        self.store.close()

    def test_all_three_formats_render(self):
        for fmt in ("json", "xlsx", "txt"):
            with self.subTest(fmt=fmt):
                body, suffix, ctype = reporting.render(self.doc, fmt)
                self.assertIsInstance(body, bytes)
                self.assertGreater(len(body), 0)
                self.assertEqual(suffix, fmt)
                self.assertIn("/", ctype)

    def test_unsupported_format_is_rejected_with_the_valid_list(self):
        with self.assertRaises(ValueError) as ctx:
            reporting.render(self.doc, "pdf")
        self.assertIn("json", str(ctx.exception))

    def test_format_choice_is_case_insensitive(self):
        for fmt in ("JSON", "Xlsx", "TXT"):
            with self.subTest(fmt=fmt):
                body, _, _ = reporting.render(self.doc, fmt)
                self.assertGreater(len(body), 0)

    def test_json_is_the_document(self):
        body, _, _ = reporting.render(self.doc, "json")
        parsed = json.loads(body)
        self.assertEqual(parsed["report"]["kind"], "summary")
        self.assertEqual(parsed["section_names"], self.doc["section_names"])


class TestTextReport(unittest.TestCase):
    def setUp(self):
        self.store = make_store()
        self.doc = reporting.build_document("summary", self.store, window="all")
        self.text = reporting.render(self.doc, "txt")[0].decode()

    def tearDown(self):
        self.store.close()

    def test_mentions_every_section(self):
        for name in self.doc["section_names"]:
            self.assertIn(name.upper(), self.text,
                          f"section {name} is missing from the text report")

    def test_states_when_and_what(self):
        meta = self.doc["report"]
        self.assertIn(meta["generated_at"], self.text)
        self.assertIn("all", self.text)

    def test_says_so_when_a_section_is_empty(self):
        """An empty table must not read as a rendering failure."""
        doc = reporting.build_document("incidents", self.store, window="all")
        text = reporting.render(doc, "txt")[0].decode()
        self.assertIn("(nothing to report)", text)

    def test_is_paste_safe(self):
        """No colour codes or markup: it may go into a plain-text ticket."""
        self.assertNotIn("\x1b[", self.text)
        self.assertNotIn("<span", self.text)
        self.assertNotIn("&nbsp;", self.text)


class TestXlsx(unittest.TestCase):
    def setUp(self):
        self.store = make_store()
        self.sid = self.store.scan_id_for_test  # type: ignore[attr-defined]
        self.doc = reporting.build_document("summary", self.store, window="all")
        self.body = reporting.render(self.doc, "xlsx")[0]

    def tearDown(self):
        self.store.close()

    def test_is_a_zip(self):
        self.assertEqual(self.body[:2], b"PK")
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            self.assertIsNone(z.testzip())

    def test_has_the_required_parts(self):
        """A workbook missing any of these will not open."""
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            names = set(z.namelist())
        for required in (
            "[Content_Types].xml",
            "_rels/.rels",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/styles.xml",
        ):
            self.assertIn(required, names, f"xlsx is missing {required}")

    def test_every_part_is_well_formed_xml(self):
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            for name in z.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    with self.subTest(part=name):
                        minidom.parseString(z.read(name))

    def test_one_sheet_per_section_plus_a_cover(self):
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            wb = z.read("xl/workbook.xml").decode()
            n_sheets = len([n for n in z.namelist()
                            if n.startswith("xl/worksheets/sheet")])
        self.assertIn("about", wb)
        self.assertEqual(n_sheets, len(self.doc["section_names"]) + 1)

    def test_every_section_name_is_a_sheet(self):
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            wb = z.read("xl/workbook.xml").decode()
        for name in self.doc["section_names"]:
            self.assertIn(f'name="{name}"', wb,
                          f"section {name} has no sheet")

    def test_has_a_named_default_style(self):
        """Excel warns on a workbook with no default cell style."""
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            styles = z.read("xl/styles.xml").decode()
        self.assertIn("<cellStyles", styles)
        self.assertIn('name="Normal"', styles)

    def test_relationships_resolve(self):
        """Every worksheet relationship must point at a part that exists."""
        with zipfile.ZipFile(io.BytesIO(self.body)) as z:
            names = set(z.namelist())
            rels = z.read("xl/_rels/workbook.xml.rels").decode()
        import re
        for target in re.findall(r'Target="([^"]+)"', rels):
            self.assertIn(f"xl/{target}", names,
                          f"relationship target {target} has no part")

    def test_special_characters_are_escaped(self):
        """A value containing <, >, or & would corrupt the XML otherwise."""
        store = make_store()
        doc = reporting.build_document("summary", store, window="all")
        doc["sections"][0]["rows"][0]["note"] = "a < b & c > d"
        body = reporting.render(doc, "xlsx")[0]
        with zipfile.ZipFile(io.BytesIO(body)) as z:
            for name in z.namelist():
                if name.endswith(".xml"):
                    minidom.parseString(z.read(name))  # raises if malformed
        store.close()

    def test_reads_in_a_real_spreadsheet_library(self):
        """Skipped when openpyxl is absent: the project has no dependencies."""
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        import warnings
        with warnings.catch_warnings():
            # A default-style warning means the file is subtly wrong.
            warnings.simplefilter("error", UserWarning)
            wb = openpyxl.load_workbook(io.BytesIO(self.body))
        self.assertIn("about", wb.sheetnames)
        for name in self.doc["section_names"]:
            self.assertIn(name, wb.sheetnames)


class TestWriteToFile(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.store = make_store()
        self.doc = reporting.build_document("summary", self.store, window="7d")

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def test_writes_a_file_and_describes_it(self):
        info = reporting.render_to_file(self.doc, "xlsx", self.tmp.name)
        self.assertTrue(info["name"].endswith(".xlsx"))
        self.assertGreater(info["size"], 0)
        self.assertEqual(info["kind"], "summary")
        self.assertEqual(info["window"], "7d")

    def test_filename_encodes_kind_window_and_time(self):
        """A folder of reports must be readable without opening anything."""
        info = reporting.render_to_file(self.doc, "txt", self.tmp.name)
        self.assertIn("summary", info["name"])
        self.assertIn("7d", info["name"])

    def test_written_file_matches_what_was_rendered(self):
        info = reporting.render_to_file(self.doc, "json", self.tmp.name)
        with open(info["path"], "rb") as fh:
            self.assertEqual(json.load(fh)["report"]["kind"], "summary")

    def test_creates_the_directory(self):
        import os
        target = os.path.join(self.tmp.name, "nested", "deeper")
        reporting.render_to_file(self.doc, "txt", target)
        self.assertTrue(os.path.isdir(target))

    def test_rejects_a_bad_format_without_writing_anything(self):
        import os
        before = set(os.listdir(self.tmp.name))
        with self.assertRaises(ValueError):
            reporting.render_to_file(self.doc, "exe", self.tmp.name)
        self.assertEqual(set(os.listdir(self.tmp.name)), before)


if __name__ == "__main__":
    unittest.main()
